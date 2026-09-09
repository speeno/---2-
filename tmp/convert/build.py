from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
D=Document(); s=D.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.7); s.left_margin=s.right_margin=Cm(1.6)
normal=D.styles['Normal']; normal.font.name='AppleMyungjo'; normal.font.size=Pt(11)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'AppleMyungjo')
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.2
p=D.add_paragraph('342'); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
T=D.add_table(rows=0, cols=2); T.style='Table Grid'; T.autofit=False; T.columns[0].width=Cm(2.2); T.columns[1].width=Cm(15.6); T.alignment=WD_TABLE_ALIGNMENT.CENTER
rows=[('十六世', [('女 賢淑  현숙','載桓　炳鍾　見上 六七\n一九六七年丁未　月　日生\n夫 鄭學然 進永人'),('女 京心  경심','載桓　炳鍾　見上 六七\n一九七二年壬子 一月十五日生'),('澤仁  택인','載桓　得鍾　見上 六七\n一九六三年癸卯 四月二十九日生\n配 許順子 金海人\n父 성학\n一九六三年癸卯 十月五日生')]),('十七世',[('子 東赫  동혁','一九九五年乙亥 四月二十三日生'),('女 可姬  가희','一九九〇年庚午 十一月二十日生'),('女 多喜  다희','一九九四年甲')]),('十八世',[]),('十九世',[]),('二十世',[]),('二十一世',[])]
for gen, people in rows:
 c=T.add_row().cells; c[0].width=Cm(2.2); c[1].width=Cm(15.6); c[0].text=gen
 c[0].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 for i,(name,body) in enumerate(people):
  p=c[1].paragraphs[0] if i==0 else c[1].add_paragraph()
  p.paragraph_format.space_before=Pt(4 if i else 0)
  r=p.add_run(name); r.bold=True; r.font.size=Pt(12)
  p=c[1].add_paragraph(body)
 if not people: c[1].paragraphs[0].paragraph_format.space_after=Pt(13)
 for cell in c:
  for p in cell.paragraphs:
   for run in p.runs:
    run.font.name='AppleMyungjo'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'AppleMyungjo')
D.save('output/전주이씨_혜령군2_400페이지_342쪽.docx')
